"""Rebuild the README images from the live system (engine running, demo data loaded).

  .venv/bin/python docs/img/make_screenshots.py [sample_input.xlsx]

parameters.png (+ window.png when a sample input file is given): the input window (ImageMagick 'import');
stock.png / shipments.png: sample rows of the status spreadsheet rendered by LibreOffice headless.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = Path(__file__).resolve().parents[2]
IMG = ROOT / "docs" / "img"
XLSX = ROOT / "data" / "out" / "warehouse_status.xlsx"
PY = sys.executable
HEAD = PatternFill("solid", fgColor="1F4E78")
FRESH = PatternFill("solid", fgColor="E2F0D9")
THIN = Side(style="thin", color="BFBFBF")


def render(df: pd.DataFrame, name: str, fresh_col: str | None, tmp: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF", name="Calibri", size=11), HEAD
        c.alignment = Alignment(vertical="center")
    for row in df.itertuples(index=False):
        ws.append([("" if pd.isna(v) else v) for v in row])
        if fresh_col and row[df.columns.get_loc(fresh_col)] == "yes":
            for c in ws[ws.max_row]:
                c.fill = FRESH
    for col in ws.columns:
        for c in col:
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if c.row > 1:
                c.font = Font(name="Calibri", size=11)
            if hasattr(c.value, "strftime"):
                c.number_format = "dd/mm hh:mm"
        width = max(len(c.value.strftime("%d/%m %H:%M")) if hasattr(c.value, "strftime") else len(str(c.value or ""))
                    for c in col)
        ws.column_dimensions[col[0].column_letter].width = width + 3
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    for m in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, m, 0.1)
    src = tmp / f"{name}.xlsx"
    wb.save(src)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/lo", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp), str(src)], check=True, capture_output=True)
    subprocess.run(["pdftoppm", "-png", "-r", "120", "-singlefile", str(tmp / f"{name}.pdf"), str(tmp / name)], check=True)
    subprocess.run(["convert", str(tmp / f"{name}.png"), "-trim", "-bordercolor", "white", "-border", "10",
                    str(IMG / f"{name}.png")], check=True)


def main() -> None:
    IMG.mkdir(parents=True, exist_ok=True)
    if len(sys.argv) > 1:
        subprocess.run([PY, str(ROOT / "desk" / "warehouse_window.py"), "--screenshot", str(IMG / "window.png"),
                        sys.argv[1]], check=True, timeout=60)
    subprocess.run([PY, str(ROOT / "desk" / "warehouse_window.py"), "--screenshot", str(IMG / "parameters.png"),
                    "--expanded"], check=True, timeout=60)
    sheets = pd.read_excel(XLSX, sheet_name=None)
    s = sheets["Stock summary"]
    pick = pd.concat([s[s["Fresh"] == "yes"].head(4),
                      s[(s["Fresh"] != "yes") & (s["Packaging"] == "pallet")].sample(4, random_state=1),
                      s[(s["Fresh"] != "yes") & (s["Packaging"] == "small")].sample(3, random_state=2),
                      s[(s["Fresh"] != "yes") & (s["Packaging"] == "parcel")].sample(3, random_state=3)])
    pick = pick.drop(columns=["Status"]).fillna("")
    sh = sheets["Shipments"]
    ship = pd.concat([sh[sh["Status"].str.startswith("released")].head(5),
                      sh[sh["Status"].str.startswith("in stock")].sample(7, random_state=4).sort_values("Ships")])
    ship = ship[["Ships", "Order", "Customer", "SKU", "Qty", "Truck", "Dock", "From case", "SHIP slot", "Status"]]
    with tempfile.TemporaryDirectory() as t:
        render(pick, "stock", "Fresh", Path(t))
        render(ship.fillna(""), "shipments", None, Path(t))
    print("images in", IMG)


if __name__ == "__main__":
    main()
