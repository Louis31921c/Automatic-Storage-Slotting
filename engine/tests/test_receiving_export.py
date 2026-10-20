import time
from datetime import timedelta

import pandas as pd
from openpyxl import load_workbook

from app.db import session_scope, utcnow
from app.export_xlsx import build_workbook
from app.settings import settings
from conftest import ingest, pallet_line, units


def dn(client, tmp_path, rows, name="dn.xlsx"):
    f = tmp_path / name
    pd.DataFrame(rows, columns=["DDT Number", "PO Number", "Line", "Delivered Qty", "Delivery Date"]).to_excel(
        f, index=False)
    r = client.post("/ingest/file", json={"event": "modified", "path": str(f)})
    assert r.status_code == 200, r.text
    return r.json()


def sheet(name):
    with session_scope() as s:
        ws = build_workbook(s)[name]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:]]


def test_po_only_is_ordered_not_in_stock(client):
    ingest(client, [pallet_line(qty=2)])
    assert [u["state"] for u in units(client, "PO-1-10")] == ["PENDING", "PENDING"]
    o = sheet("Ordered")
    assert o[0]["PO"] == "PO-1" and o[0]["Qty outstanding"] == 2 and o[0]["Qty delivered"] == 0
    assert sheet("In stock") == []


def test_partial_delivery_books_cases(client, tmp_path):
    ingest(client, [pallet_line(qty=2)])
    r = dn(client, tmp_path, [["DDT-1", "PO-1", "10", 1, "18/10/2026"]])
    assert r["inserted"] == 1
    states = sorted(u["state"] for u in units(client, "PO-1-10"))
    assert states == ["PENDING", "PLACED"]
    o = sheet("Ordered")
    assert o[0]["Qty delivered"] == 1 and o[0]["Qty outstanding"] == 1
    stock = sheet("In stock")
    assert len(stock) == 1 and stock[0]["Case #"].startswith("A1S") and stock[0]["Status"] == "stored"
    assert stock[0]["Fresh"] == "yes" and stock[0]["Building"] == "B2"
    # same DN again: unchanged; second DN completes the line
    again = dn(client, tmp_path, [["DDT-1", "PO-1", "10", 1, "18/10/2026"]])     # rewritten file, same row
    assert (again["skipped_file"] or again["unchanged"] == 1) and again["inserted"] == 0
    dn(client, tmp_path, [["DDT-2", "PO-1", "10", 1, "19/10/2026"]], name="dn2.xlsx")
    assert sheet("Ordered") == [] and len(sheet("In stock")) == 2


def test_oversize_received_shows_extra(client, tmp_path):
    ingest(client, [pallet_line(h=2500)])
    dn(client, tmp_path, [["DDT-1", "PO-1", "10", None, None]])                # blank qty = whole line
    stock = sheet("In stock")
    assert stock[0]["Case #"] == "EXTRA" and stock[0]["Status"].startswith("EXTRA")


def test_tablet_mode_receives_without_booking(client, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "auto_place_on_receipt", False)
    ingest(client, [pallet_line()])
    dn(client, tmp_path, [["DDT-1", "PO-1", "10", 1, None]])
    assert units(client, "PO-1-10")[0]["state"] == "PENDING"
    assert sheet("In stock")[0]["Status"] == "waiting for a case"


def test_stale_receipt_not_fresh(client, tmp_path):
    ingest(client, [pallet_line()])
    dn(client, tmp_path, [["DDT-1", "PO-1", "10", 1, None]])
    with session_scope() as s:
        assert build_workbook(s, now=utcnow() + timedelta(days=2))["In stock"]["M2"].value in (None, "")


def test_shipments_sheet(client, tmp_path):
    ingest(client, [pallet_line(qty=1)])
    dn(client, tmp_path, [["DDT-1", "PO-1", "10", 1, None]])
    case = sheet("In stock")[0]["Case #"]
    f = tmp_path / "cpo.xlsx"
    ship = (utcnow() + timedelta(days=4)).replace(microsecond=0)
    pd.DataFrame([["C-1", 1, "Garage", "SKU-A", 2, ship, None], ["C-2", 1, "Shop", "SKU-Z", 1, None, None]],
                 columns=["Order Number", "Line", "Customer", "SKU", "Qty", "Ship Date", "Truck"]).to_excel(f, index=False)
    client.post("/ingest/file", json={"event": "modified", "path": str(f)})
    rows = sheet("Shipments")
    c1 = [r for r in rows if r["Order"] == "C-1"][0]
    assert c1["From case"] == case and "short 1" in c1["Status"] and c1["Leaves case at"] is not None
    c2 = [r for r in rows if r["Order"] == "C-2"][0]
    assert c2["Status"] == "no ship time yet"
    assert rows[0]["Order"] == "C-1"                                             # dated first


def test_background_export_file_and_download(client, tmp_path):
    ingest(client, [pallet_line()])
    for _ in range(50):
        if settings.export_path.exists():
            break
        time.sleep(0.1)
    wb = load_workbook(settings.export_path)
    assert wb.sheetnames == ["Ordered", "In stock", "Stock summary", "Shipments"]
    r = client.get("/export.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_delivery_note_before_po(client, tmp_path):
    r = dn(client, tmp_path, [["DDT-9", "PO-1", "10", 2, None]])
    assert r["inserted"] == 1 and r["rejected"] == 0
    ingest(client, [pallet_line(qty=2)])
    assert [u["state"] for u in units(client, "PO-1-10")] == ["PLACED", "PLACED"]


def test_stock_summary_sheet(client, tmp_path):
    ingest(client, [pallet_line(qty=2, fragility=1)])
    dn(client, tmp_path, [["DDT-1", "PO-1", "10", 2, None]])
    full, short = sheet("In stock"), sheet("Stock summary")
    assert list(short[0]) == ["Case #", "PO", "Qty", "Status", "Fresh", "Received", "SKU", "Description",
                              "Packaging", "kg/unit", "Fragile"]
    assert [(r["Case #"], r["Qty"], r["kg/unit"]) for r in short] == \
           [(r["Case #"], r["Qty"], r["kg/unit (incl. pack)"]) for r in full]
    assert short[0]["Fragile"] == "fragile"
