"""Status spreadsheet for the team (read-only view of the DB, regenerated on every state change).

Separate exporter, never in the ingest path (consignes: "write-back via a separate exporter").
Sheets:
  Ordered     supplier PO lines with units not delivered yet
  In stock    units in the warehouse (stored + freshly received), with their case #
  Stock summary  same rows as In stock, short column set
  Shipments   open customer order lines, allocated units, source case, SHIP slot, truck
Times shown in Europe/Rome. Written atomically (tmp file + rename).

  .venv/bin/python -m app.export_xlsx            (from engine/)
"""
import logging
import os
import re
import tempfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import case, func, select
from sqlalchemy.orm import Session, selectinload

from .db import utcnow
from .models import (CANCELLED, EXTRA, PENDING, PLACED, RESERVED, SHIPPED, STAGED, Article, Case, OrderLine, PoLine,
                     Truck)
from .settings import settings
from .tables import OWN_EXPORT_TITLE

log = logging.getLogger("export")
FRESH = PatternFill("solid", fgColor="E2F0D9")
EXTRA_FILL = PatternFill("solid", fgColor="FCE4D6")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
FRAG = {0: "", 1: "fragile", 2: "very fragile"}


def local(dt: datetime | None) -> datetime | None:
    """naive UTC -> naive Europe/Rome (Excel has no time zones)."""
    if dt is None:
        return None
    return dt.replace(tzinfo=UTC, microsecond=0).astimezone(ZoneInfo(settings.tz)).replace(tzinfo=None)


def case_cols(c: Case | None, slot: str | None) -> list:
    """Case #, building, aisle, shelf, level, position."""
    if c is None:
        return [slot or "", None, None, None, None, None]
    m = re.match(r"^A(\d+)S(\d+)L(\d+)P(\d+)$", c.id)
    if m:
        return [c.id, f"B{c.building}", *(int(g) for g in m.groups())]
    return [c.id, f"B{c.building}" if c.storage_type != "ship" else "Shipping area", None, None, c.level, None]


def dims_mm(line: PoLine) -> list:
    return [line.length_mm, line.width_mm, line.height_mm, line.weight_kg]


# ---------------------------------------------------------------- sheet builders
def ordered_rows(s: Session) -> list[list]:
    rows = []
    for line in s.scalars(select(PoLine).where(PoLine.cancelled.is_(False))
                          .options(selectinload(PoLine.units), selectinload(PoLine.po))
                          .order_by(PoLine.expected_date, PoLine.po_number, PoLine.line_ref)):
        active = [a for a in line.units if a.state != CANCELLED]
        delivered = sum(1 for a in active if a.delivered_at is not None)
        outstanding = len(active) - delivered
        if outstanding <= 0:
            continue
        rows.append([line.po_number, line.line_ref, line.po.supplier, line.sku, line.description, line.qty_units,
                     delivered, outstanding, line.packaging, *dims_mm(line), line.dims_source or "missing",
                     FRAG.get(line.fragility, ""), local(line.expected_date)])
    return rows


def stock_rows(s: Session, now: datetime) -> list[tuple[list, str]]:
    """(row, tone) per case + PO line (+ customer allocation): units physically here, placed or reserved,
    or received but not stored yet (EXTRA / no dims). Grouped in SQL: one row per bin line, not per box."""
    fresh_after = now - timedelta(hours=settings.fresh_hours)
    in_case = [PLACED, RESERVED]
    vol = Article.length_m * Article.width_m * Article.height_m
    units_in = dict(s.execute(select(Article.case_id, func.count()).where(Article.state.in_(in_case))
                              .group_by(Article.case_id)).all())
    used_m3 = dict(s.execute(select(Article.case_id, func.sum(vol)).where(Article.state.in_(in_case))
                             .group_by(Article.case_id)).all())
    arrived = func.max(func.coalesce(Article.delivered_at, Article.placed_at, Article.received_at))
    q = (select(Article.case_id, Article.po_line_id, Article.state, Article.slot, Article.order_line_id,
                Article.ship_at, func.count(), arrived, func.max(Article.length_m), func.max(Article.width_m),
                func.max(Article.height_m), func.max(Article.weight_kg), func.max(Article.fragility),
                func.sum(case((Article.length_m.is_(None), 1), else_=0)))
         .where(Article.state.in_([PLACED, RESERVED, PENDING]))
         .where((Article.state != PENDING) | Article.delivered_at.is_not(None))
         .group_by(Article.case_id, Article.po_line_id, Article.state, Article.slot, Article.order_line_id,
                   Article.ship_at))
    groups = s.execute(q).all()
    lines = {ln.id: ln for ln in s.scalars(select(PoLine).where(PoLine.id.in_({g[1] for g in groups})))}
    cases = {c.id: c for c in s.scalars(select(Case).where(Case.id.in_({g[0] for g in groups if g[0]})))}
    orders = dict(s.execute(select(OrderLine.id, OrderLine.order_number)
                            .where(OrderLine.id.in_({g[4] for g in groups if g[4]}))).all())
    out = []
    for cid, line_id, state, slot, ol_id, ship_at, qty, last_in, L, Wd, H, kg, frag, no_dims in groups:
        line, c = lines[line_id], cases.get(cid)
        if state == PLACED:
            status = "stored"
        elif state == RESERVED:
            status = "case reserved, not confirmed"
        elif slot == EXTRA:
            status = "EXTRA: fits no case, manager to decide"
        elif no_dims:
            status = "dimensions needed"
        else:
            status = "waiting for a case"
        fresh = last_in is not None and last_in >= fresh_after
        cc = case_cols(c, slot if state == PENDING else None)
        fill = round(100 * (used_m3.get(cid) or 0) / c.volume_m3) if c is not None and c.volume_m3 else None
        row = [cc[0], line.po_number, line.line_ref, qty, *cc[1:], units_in.get(cid) if c else None, fill, status,
               "yes" if fresh else "", local(last_in), line.sku, line.description, line.qty_units, line.packaging,
               L, Wd, H, kg, FRAG.get(frag or 0, ""), orders.get(ol_id, ""), local(ship_at)]
        tone = "extra" if slot == EXTRA or no_dims else "fresh" if fresh else ""
        out.append((row, tone, (0 if fresh else 1, cid is None, _nat(cid or ""), line.po_number, line.line_ref)))
    out.sort(key=lambda t: t[2])
    return [(r, tone) for r, tone, _ in out]


def shipment_rows(s: Session) -> list[tuple[list, str]]:
    out = []
    for ol in s.scalars(select(OrderLine).where(OrderLine.cancelled.is_(False), OrderLine.fulfilled.is_(False))):
        truck = s.get(Truck, ol.truck_ref) if ol.truck_ref else None
        ship_at = truck.departure_at if (truck and not truck.cancelled and truck.departure_at) else ol.ship_at
        units = list(s.scalars(select(Article).where(Article.order_line_id == ol.id, Article.state != SHIPPED)
                               .order_by(Article.id)))
        base = [local(ship_at), local(ship_at - timedelta(hours=settings.pre_ship_release_h)) if ship_at else None,
                ol.order_number, ol.line_ref, ol.customer, ol.sku, ol.qty_units, ol.truck_ref or "",
                truck.carrier if truck else "", (truck.dock or "").replace("Dock_", "") if truck else ""]
        if not units:
            status = "no ship time yet" if ship_at is None else "no stock allocated"
            out.append((base + ["", "", "", "", status], "warn", ship_at))
            continue
        short = ol.qty_units - len(units)
        for a in units:
            if a.state == STAGED:
                status = "at shipping area, to load" if a.picked_at else "released, to pick"
                src, ship_slot = a.staged_from or "", a.slot or ""
            else:
                status, src, ship_slot = "in stock, allocated", a.case_id or "", ""
            if short > 0:
                status += f" (short {short})"
            out.append((base + [a.code, src, ship_slot, a.state, status], "warn" if short > 0 else "", ship_at))
    out.sort(key=lambda t: (t[2] is None, t[2] or datetime.max, t[0][2], t[0][3]))
    return [(r, tone) for r, tone, _ in out]


def _nat(x: str) -> tuple:
    return tuple(int(t) if t.isdigit() else t for t in re.findall(r"\d+|\D+", x))


# ---------------------------------------------------------------- workbook
HEADERS = {
    "Ordered": ["PO", "Line", "Supplier", "SKU", "Description", "Qty ordered", "Qty delivered", "Qty outstanding",
                "Packaging", "Length mm", "Width mm", "Height mm", "Weight kg", "Dims source", "Fragile",
                "Expected"],
    "In stock": ["Case #", "PO", "Line", "Qty", "Building", "Aisle", "Shelf", "Level", "Pos", "Units in case",
                 "Case fill %", "Status", "Fresh", "Received", "SKU", "Description", "Line qty", "Packaging",
                 "L m", "W m", "H m", "kg/unit (incl. pack)",
                 "Fragile", "Customer order", "Ships"],
    "Shipments": ["Ships", "Leaves case at", "Order", "Line", "Customer", "SKU", "Qty", "Truck", "Carrier", "Dock",
                  "Unit", "From case", "SHIP slot", "State", "Status"],
}
# short version of In stock: same rows, fewer columns (header -> In stock header it comes from)
SUMMARY = {"Case #": "Case #", "PO": "PO", "Qty": "Qty", "Status": "Status", "Fresh": "Fresh",
           "Received": "Received", "SKU": "SKU", "Description": "Description", "Packaging": "Packaging",
           "kg/unit": "kg/unit (incl. pack)", "Fragile": "Fragile"}
HEADERS["Stock summary"] = list(SUMMARY)
DATE_COLS = {"Ordered": {"Expected"}, "In stock": {"Received", "Ships"}, "Stock summary": {"Received"},
             "Shipments": {"Ships", "Leaves case at"}}


def summary_rows(stock: list[tuple[list, str]]) -> list[tuple[list, str]]:
    src = HEADERS["In stock"]
    cols = [src.index(h) for h in SUMMARY.values()]
    return [([row[i] for i in cols], tone) for row, tone in stock]


def _sheet(wb: Workbook, title: str, rows: list[tuple[list, str]], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    head = HEADERS[title]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    fills = {"fresh": FRESH, "extra": EXTRA_FILL, "warn": WARN_FILL}
    for row, tone in rows:
        ws.append(row)
        if tone in fills:
            for c in ws[ws.max_row]:
                c.fill = fills[tone]
    dates = {i + 1 for i, h in enumerate(head) if h in DATE_COLS[title]}
    for col in dates:
        for (c,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            c.number_format = "dd/mm/yyyy hh:mm"
    for i, h in enumerate(head, start=1):
        width = max([len(h)] + [len(str(r[i - 1])) for r, _ in rows[:500] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 42) if i not in dates else 17
    ws.freeze_panes = {"In stock": "D2", "Stock summary": "C2"}.get(title, "A2")   # case #, PO stay visible
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def build_workbook(s: Session, now: datetime | None = None) -> Workbook:
    now = now or utcnow()
    wb = Workbook()
    _sheet(wb, "Ordered", [(r, "") for r in ordered_rows(s)], first=True)
    stock = stock_rows(s, now)
    _sheet(wb, "In stock", stock)
    _sheet(wb, "Stock summary", summary_rows(stock))
    _sheet(wb, "Shipments", shipment_rows(s))
    wb.properties.title = OWN_EXPORT_TITLE
    wb.properties.description = f"Generated {local(now):%d/%m/%Y %H:%M} (Europe/Rome). Read-only view; edit the inbox files instead."
    return wb


def write_export(s: Session, path: Path | None = None) -> Path:
    path = path or settings.export_path
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(s)
    fd, tmp = tempfile.mkstemp(prefix=".tmp_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.chmod(tmp, 0o644)                # mkstemp creates 0600; the team must be able to read it
        os.replace(tmp, path)              # atomic: readers never see a half-written file
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path


if __name__ == "__main__":
    from .db import init_db, session_scope
    init_db()
    with session_scope() as s:
        print(write_export(s))
